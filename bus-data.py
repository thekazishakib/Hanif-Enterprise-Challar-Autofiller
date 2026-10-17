"""Pricing, Excel inputs, signatures and durable shared challan numbering."""
from __future__ import annotations
import base64
from datetime import date,datetime,time,timedelta
from decimal import Decimal,InvalidOperation,ROUND_HALF_UP
from io import BytesIO
from pathlib import Path
import sqlite3

SEATS=[f'{a}-{b}' for a in 'ABCDEFGHI' for b in range(1,5)]
FIELDS=['Challan No','Vehicle No','Coach No','Permit No','Driver','Driver Mobile','Guide','Guide Mobile','Checker Name','Route','Date','Time','Coaster No','Coaster Date','Coaster Time','Coaster Fare','Book Cash','Counter Deposit','Refund Deposit','Fuel Deposit','Wash Deposit','B5 Deposit','Sohoj Cash']
CASH=FIELDS[16:]
SHARED=['Vehicle No','Coach No','Permit No','Guide','Guide Mobile','Checker Name','Route']
BN=str.maketrans('0123456789','০১২৩৪৫৬৭৮৯'); EN=str.maketrans('০১২৩৪৫৬৭৮৯','0123456789')

class InputError(ValueError): pass

def text(v): return '' if v is None else str(v).strip()

def amount(v,label,default=None):
    if not text(v) and default is not None: return Decimal(default)
    try:
        n=Decimal(text(v).translate(EN))
        if not n.is_finite() or n<0 or n>Decimal('999999999.99') or n!=n.quantize(Decimal('.01')): raise ValueError()
    except (InvalidOperation,ValueError): raise InputError(f'{label}: enter a non-negative amount with at most two decimals.') from None
    return n

def parsed_date(v,label):
    if isinstance(v,datetime): return v.date().isoformat()
    if isinstance(v,date): return v.isoformat()
    try: return date.fromisoformat(text(v).translate(EN)).isoformat()
    except ValueError: raise InputError(f'{label}: use an Excel date or YYYY-MM-DD.') from None

def parsed_time(v,label):
    if isinstance(v,datetime): v=v.time()
    if isinstance(v,time): return v.strftime('%H:%M')
    if isinstance(v,(int,float)) and not isinstance(v,bool) and 0<=v<1:
        return (datetime(2000,1,1)+timedelta(days=v)).strftime('%H:%M')
    try: return time.fromisoformat(text(v).translate(EN)).strftime('%H:%M')
    except ValueError: raise InputError(f'{label}: use an Excel time or HH:MM.') from None

def value(cell,sheet):
    if cell.data_type in ('f','e'): raise InputError(f'{sheet}!{cell.coordinate}: enter a value, not a formula or error.')
    return cell.value

def read_catalog(sheet,headers):
    if [text(c.value) for c in sheet[4]][:len(headers)]!=headers:
        raise InputError(f'{sheet.title}: keep the original row 4 headers.')
    out={}
    for row in sheet.iter_rows(min_row=5,max_row=104,max_col=len(headers)):
        if not any(text(c.value) for c in row): continue
        cells=[value(c,sheet.title) for c in row]; name=text(cells[0])
        if not name: raise InputError(f'{sheet.title} row {row[0].row}: name is required.')
        if name in out: raise InputError(f'{sheet.title}: duplicate name {name}.')
        out[name]=dict(zip(headers,cells))
    for row in sheet.iter_rows(min_row=105,max_col=len(headers)):
        if any(text(c.value) for c in row): raise InputError(f'{sheet.title}: use catalog rows 5-104 (100 entries).')
    return out

def catalogs(wb):
    routes=read_catalog(wb['Routes'],['Route','From','To','Fare BDT','Route EN','From EN','To EN','Route No'])
    for name,r in routes.items():
        if not text(r['From']) or not text(r['To']): raise InputError(f'Routes: {name} needs From and To.')
        r['Fare BDT']=amount(r['Fare BDT'],f'Routes: {name}')
    counters=read_catalog(wb['Counters'],['Counter','Commission Mode','Rate','Signatory','Signature File','Counter EN','Signatory EN'])
    for name,c in counters.items():
        mode=text(c['Commission Mode']) or 'None'
        c['Commission Mode']=mode
        if mode not in ('None','Fixed','Percent'): raise InputError(f'Counters: {name}, choose None, Fixed or Percent.')
        rate=amount(c['Rate'],f'Counters: {name}, Rate',default=0 if mode=='None' else None)
        if mode=='None' and rate!=0: raise InputError(f'Counters: {name}, None must have a zero or blank rate.')
        if mode=='Percent' and rate>100: raise InputError(f'Counters: {name}, percent cannot exceed 100.')
        c['Rate']=rate
    return routes,counters

def signature_image(raw,base):
    if not text(raw): return ''
    from PIL import Image
    path=Path(text(raw)); path=path if path.is_absolute() else base/path
    if path.stat().st_size>2_000_000: raise InputError('Signature image must be smaller than 2 MB.')
    data=path.read_bytes()
    try:
        with Image.open(BytesIO(data)) as im:
            kind=im.format
            if kind not in ('PNG','JPEG') or max(im.size)>4000: raise InputError('Use a PNG/JPEG signature up to 4000 pixels.')
            im.verify()
    except (OSError,ValueError) as exc: raise InputError(f'Cannot read signature image: {path.name}') from exc
    return 'data:image/'+('png' if kind=='PNG' else 'jpeg')+';base64,'+base64.b64encode(data).decode('ascii')

SIGNATURE_ROLES=('booking_master','manager','checker')
def role_signatures(config,base):
    """Fixed signature images for Booking master, Manager and Checker.
    Use settings.json (<role>_signature) or drop <role-with-dashes>.png/.jpg into the signatures folder."""
    found={}
    for role in SIGNATURE_ROLES:
        configured=text(config.get(role+'_signature'))
        path=None
        if configured:
            path=Path(configured); path=path if path.is_absolute() else base/path
            if not path.is_file(): raise InputError(f'Signature file for {role.replace("_"," ")} not found: {configured}')
        else:
            for ext in ('png','jpg','jpeg'):
                cand=base/'signatures'/f'{role.replace("_","-")}.{ext}'
                if cand.is_file(): path=cand;break
        if path: found[role]=signature_image(str(path),base)
    return found

def read_details(sheet,routes):
    raw=[{},{}]; seen=set()
    for row in sheet.iter_rows(min_row=5,max_col=3):
        key=text(row[0].value)
        if not key: continue
        if key not in FIELDS or key in seen: raise InputError(f'Trip Details: unknown/duplicate field {key}.')
        seen.add(key)
        if key=='Challan No':
            raw[0][key]=raw[1][key]=''; continue
        raw[0][key]=value(row[1],sheet.title)
        raw[1][key]=raw[0][key] if key in SHARED else value(row[2],sheet.title)
    if seen!=set(FIELDS): raise InputError('Use the updated Trip Details fields; some fields are missing.')
    changed_driver=bool(text(raw[1]['Driver'])) and text(raw[1]['Driver'])!=text(raw[0]['Driver'])
    if not text(raw[1]['Driver']):raw[1]['Driver']=raw[0]['Driver']
    if not text(raw[1]['Driver Mobile']) and not changed_driver:raw[1]['Driver Mobile']=raw[0]['Driver Mobile']
    result=[]
    for i,r in enumerate(raw):
        name=['Main','Secondary'][i]; d={k:text(v) for k,v in r.items()}
        if not d['Vehicle No']: raise InputError('Main: Vehicle No is required.')
        route=routes.get(d['Route'])
        if not route: raise InputError('Select a configured route on Trip Details.')
        d['Date']=parsed_date(r['Date'],f'{name} date'); d['Time']=parsed_time(r['Time'],f'{name} time')
        d['From']=text(route['From' if i==0 else 'To']);d['To']=text(route['To' if i==0 else 'From'])
        d['From EN']=text(route['From EN' if i==0 else 'To EN']) or d['From']
        d['To EN']=text(route['To EN' if i==0 else 'From EN']) or d['To']
        d['Route EN']=text(route['Route EN']) or d['Route'];d['Route No']=text(route['Route No'])
        d['Base Fare']=route['Fare BDT']
        for key in CASH: d[key]=amount(r[key],f'{name}: {key}',default=0)
        d['Coaster Fare']=amount(r['Coaster Fare'],f'{name}: Coaster Fare',default=0)
        if d['Coaster No']:
            d['Coaster Date']=parsed_date(r['Coaster Date'],f'{name}: Coaster Date')
            d['Coaster Time']=parsed_time(r['Coaster Time'],f'{name}: Coaster Time')
        elif d['Coaster Date'] or d['Coaster Time'] or d['Coaster Fare']:
            raise InputError(f'{name}: enter Coaster No for external-coaster date/time/fare.')
        result.append(d)
    return result

def read_seats(sheet,details,counters,base,images):
    headers=['Seat No','Counter','Passengers','Fare BDT','Commission BDT','Net BDT','Signature','Fare Override']
    if [text(c.value) for c in sheet[4]][:8]!=headers: raise InputError(f'{sheet.title}: use the updated seat template.')
    result={}
    for row in sheet.iter_rows(min_row=5,max_row=40,max_col=8):
        seat=text(value(row[0],sheet.title)); name=text(value(row[1],sheet.title)); override=value(row[7],sheet.title)
        if seat not in SEATS or seat in result: raise InputError(f'{sheet.title}: invalid or duplicate seat {seat}.')
        if not name:
            if text(override): raise InputError(f'{sheet.title} {seat}: choose a counter before overriding fare.')
            result[seat]=None;continue
        if name not in counters: raise InputError(f'{sheet.title} {seat}: counter {name!r} is not configured.')
        c=counters[name]; fare=amount(override,f'{sheet.title} {seat} fare') if text(override) else details['Base Fare']
        mode=text(c['Commission Mode']); rate=c['Rate']
        commission=Decimal(0) if mode=='None' else rate if mode=='Fixed' else (fare*rate/100).quantize(Decimal('.01'),rounding=ROUND_HALF_UP)
        if commission>fare: raise InputError(f'{sheet.title} {seat}: commission exceeds fare.')
        if name not in images: images[name]=signature_image(c['Signature File'],base)
        result[seat]={'Counter':name,'Counter EN':text(c['Counter EN']) or name,'Passengers':1,'Fare BDT':fare,'Commission BDT':commission,'Net BDT':fare-commission,'Signature':text(c['Signatory']),'Signature EN':text(c['Signatory EN']) or text(c['Signatory']),'Signature Image':images[name]}
    if set(result)!=set(SEATS): raise InputError(f'{sheet.title}: keep all 36 seats.')
    for row in sheet.iter_rows(min_row=42,max_col=8):
        if text(row[0].value) in SEATS or any(text(c.value) for c in row[1:]): raise InputError(f'{sheet.title}: keep seat entries within rows 5-40.')
    return result

def load_data(path):
    from openpyxl import load_workbook
    path=Path(path); wb=load_workbook(path,data_only=False)
    try:
        for name in ['Trip Details','Main','Secondary','Routes','Counters']:
            if name not in wb.sheetnames: raise InputError(f'Missing {name}. Use the new dropdown template.')
        routes,counters=catalogs(wb); details=read_details(wb['Trip Details'],routes); images={}
        return [{'details':details[i],'seats':read_seats(wb[name],details[i],counters,path.parent,images)} for i,name in enumerate(['Main','Secondary'])]
    finally: wb.close()

def blank_data():
    d={k:Decimal(0) if k in CASH or k=='Coaster Fare' else '' for k in FIELDS}
    d.update({'From':'','To':'','Base Fare':Decimal(0)})
    return [{'details':dict(d),'seats':{s:None for s in SEATS}} for _ in range(2)]

def connect_register(path):
    path=Path(path);path.parent.mkdir(parents=True,exist_ok=True)
    con=sqlite3.connect(path,timeout=30)
    con.execute('CREATE TABLE IF NOT EXISTS challans (year INTEGER NOT NULL, serial INTEGER NOT NULL, number TEXT NOT NULL, created TEXT NOT NULL, html TEXT NOT NULL, html_en TEXT, PRIMARY KEY(year,serial))')
    if 'html_en' not in [r[1] for r in con.execute('PRAGMA table_info(challans)')]:
        con.execute('ALTER TABLE challans ADD COLUMN html_en TEXT')
    return con

def save_numbered(trips,config,render,registry):
    year=date.fromisoformat(trips[0]['details']['Date']).year
    con=connect_register(registry)
    try:
        con.execute('BEGIN IMMEDIATE')
        serial=con.execute('SELECT COALESCE(MAX(serial),0)+1 FROM challans WHERE year=?',(year,)).fetchone()[0]
        number=f'{serial:02d}/{year}'
        for trip in trips: trip['details']['Challan No']=number.translate(BN)
        html=render(trips,config,'bn'); html_en=render(trips,config,'en')
        con.execute('INSERT INTO challans(year,serial,number,created,html,html_en) VALUES(?,?,?,?,?,?)',(year,serial,number,datetime.now().isoformat(timespec='seconds'),html,html_en))
        con.commit();return number,{'bn':html,'en':html_en}
    except Exception:
        con.rollback();raise
    finally: con.close()

def list_saved(registry):
    if not Path(registry).exists(): return []
    with connect_register(registry) as con: return con.execute('SELECT number,created FROM challans ORDER BY year DESC,serial DESC').fetchall()

def saved_html(registry,number,language='bn'):
    if not Path(registry).exists(): raise InputError('No saved challans yet.')
    column='html_en' if language=='en' else 'html'
    with connect_register(registry) as con: row=con.execute(f'SELECT {column} FROM challans WHERE number=?',(number.translate(EN),)).fetchone()
    if not row: raise InputError('That challan number is not in this register.')
    if not row[0]: raise InputError('This older record has no saved English version.')
    return row[0]
