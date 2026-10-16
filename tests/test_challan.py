from concurrent.futures import ThreadPoolExecutor
from datetime import date,time
from decimal import Decimal
from pathlib import Path
import sys,tempfile,unittest
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
import bus_data as data
import challan_autofill as app
from print_forms import document

class ChallanTests(unittest.TestCase):
    def setUp(self):
        from openpyxl import load_workbook
        self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup)
        self.folder=Path(self.tmp.name);self.path=self.folder/'input.xlsx';self.registry=self.folder/'register.sqlite3'
        self.wb=load_workbook(app.BASE/'challan_input.xlsx');self.addCleanup(self.wb.close)
        s=self.wb['Trip Details']
        for cell,v in {'B6':'TEST-VEHICLE','B9':'Driver A','B10':'01000000000','B11':'Guide A','B12':'01000000001','B13':'Checker A','B14':'ঢাকা–চট্টগ্রাম','B15':date(2026,10,16),'C15':date(2026,10,17),'B16':time(9),'C16':time(12)}.items():s[cell]=v
        c=self.wb['Counters']
        for r,values in enumerate([['DIRECT','None',0,'Signer',None,'DIRECT','Signer'],['PERCENT','Percent',10,None,None,'PERCENT',None],['FIXED','Fixed',20,None,None,'FIXED',None]],5):
            for col,v in enumerate(values,1):c.cell(r,col,v)
        for r,counter in enumerate(['DIRECT','PERCENT','FIXED'],5):self.wb['Main'].cell(r,2,counter)
        self.wb['Secondary']['B5']='DIRECT'
    def save(self):self.wb.save(self.path);return self.path
    def load(self):return data.load_data(self.save())
    def generate(self):return app.generate(self.save(),self.folder/'output',registry=self.registry)
    def test_automatic_fare_commission_and_passengers(self):
        trips=self.load();s=trips[0]['seats']
        self.assertEqual([s[k]['Commission BDT'] for k in ['A-1','A-2','A-3']],[0,30,20])
        self.assertEqual(sum(i['Net BDT'] for i in s.values() if i),850)
        self.assertEqual(sum(i['Passengers'] for i in s.values() if i),3)
        self.assertIsNone(s['I-4'])
    def test_shared_route_crew_and_optional_driver_override(self):
        self.wb['Trip Details']['C9']='Driver B'
        a,b=self.load()
        for field in ['Vehicle No','Guide','Guide Mobile','Checker Name','Route']:
            self.assertEqual(a['details'][field],b['details'][field])
        self.assertEqual(b['details']['Driver'],'Driver B')
        self.assertEqual(b['details']['Driver Mobile'],'')
        self.assertEqual(a['details']['From'],b['details']['To'])
        self.assertNotEqual(a['details']['Date'],b['details']['Date'])
    def test_counter_selection_blank_and_route_changes(self):
        self.wb['Main']['B5']=None;self.wb['Routes']['D5']=400
        trips=self.load();self.assertIsNone(trips[0]['seats']['A-1'])
        self.assertEqual(trips[0]['seats']['A-2']['Fare BDT'],400)
        self.assertEqual(trips[1]['seats']['A-1']['Fare BDT'],400)
    def test_zero_override(self):
        self.wb['Main']['H5']=0
        self.assertEqual(self.load()[0]['seats']['A-1']['Net BDT'],0)
    def test_no_commission_is_default(self):
        self.wb['Counters']['B5']=None;self.wb['Counters']['C5']=None
        self.assertEqual(self.load()[0]['seats']['A-1']['Commission BDT'],0)
    def test_missing_commission_is_not_silently_zero(self):
        self.wb['Counters']['C6']=None
        with self.assertRaises(data.InputError):self.load()
    def test_unknown_counter_and_duplicate_seats(self):
        self.wb['Main']['B5']='UNKNOWN'
        with self.assertRaisesRegex(data.InputError,'not configured'):self.load()
        self.wb['Main']['B5']='DIRECT';self.wb['Main']['A6']='A-1'
        with self.assertRaisesRegex(data.InputError,'duplicate'):self.load()
    def test_invalid_commission_and_negative_fare(self):
        self.wb['Counters']['C6']=101
        with self.assertRaisesRegex(data.InputError,'100'):self.load()
        self.wb['Counters']['C6']=10;self.wb['Main']['H5']=-1
        with self.assertRaises(data.InputError):self.load()
    def test_date_time_required_independently(self):
        self.wb['Trip Details']['C15']=None
        with self.assertRaises(data.InputError):self.load()
    def test_external_coaster_and_zero_deposits(self):
        s=self.wb['Trip Details'];s['B17']='EXT-1';s['B18']=date(2026,10,16);s['B19']=time(8);s['B20']=1200
        d=self.load()[0]['details'];self.assertEqual(d['Coaster Fare'],1200)
        for k in ['Refund Deposit','Fuel Deposit','Wash Deposit']:self.assertEqual(d[k],0)
        s['B18']=None
        with self.assertRaises(data.InputError):self.load()
    def test_numbers_shared_increment_and_reprint_unchanged(self):
        first=self.generate();before=first['bn'].read_text(encoding='utf-8')
        self.assertEqual(before.count('০১/২০২৬'),2)
        self.assertEqual(first['en'].read_text(encoding='utf-8').count('01/2026'),2)
        again=app.reprint('01/2026',self.folder/'reprints',registry=self.registry)
        self.assertEqual(again['bn'].read_text(encoding='utf-8'),before)
        self.assertEqual(len(data.list_saved(self.registry)),1)
        second=self.generate();self.assertEqual(second['bn'].read_text(encoding='utf-8').count('০২/২০২৬'),2)
    def test_failed_validation_does_not_consume_number(self):
        self.wb['Trip Details']['B15']='bad'
        with self.assertRaises(data.InputError):self.generate()
        self.assertFalse(self.registry.exists())
    def test_failed_render_rolls_back(self):
        trips=self.load()
        def bad(*args):raise RuntimeError('render failure')
        with self.assertRaises(RuntimeError):data.save_numbered(trips,{},bad,self.registry)
        number,_=data.save_numbered(trips,{},document,self.registry);self.assertEqual(number,'01/2026')
    def test_concurrent_saves_get_distinct_numbers(self):
        self.save()
        def save_one(_):return data.save_numbered(data.load_data(self.path),{},document,self.registry)[0]
        with ThreadPoolExecutor(max_workers=4) as pool:numbers=list(pool.map(save_one,range(4)))
        self.assertEqual(set(numbers),{'01/2026','02/2026','03/2026','04/2026'})
    def test_year_uses_outgoing_trip_date(self):
        self.generate();self.wb['Trip Details']['B15']=date(2027,1,1)
        self.assertIn('01/2027',self.generate()['en'].read_text(encoding='utf-8'))
    def test_blank_preview_does_not_number(self):
        p=app.generate(None,self.folder,blank=True,registry=self.registry)
        self.assertFalse(self.registry.exists());self.assertEqual(set(p),{'bn','en'})
    def test_merged_checker_bilingual_headers_and_fonts(self):
        trips=self.load();config=app.load_config(app.BASE/'settings.json')
        for lang in ['bn','en']:
            html=document(trips,config,lang)
            self.assertEqual(html.count('rowspan="37"'),2)
            self.assertEqual(html.count('class="seat"'),72)
            self.assertEqual(html.count('Checker A'),2)
        en=document(trips,config,'en');self.assertIn('Times New Roman',en);self.assertIn('Head office',en);self.assertIn('Hanif Enterprise',en)
    def test_signature_image_embedded_and_source_text_escaped(self):
        from PIL import Image
        Image.new('RGB',(100,30),'white').save(self.folder/'signature.png')
        self.wb['Counters']['E5']='signature.png';self.wb['Trip Details']['B9']='<script>alert(1)</script>'
        html=document(self.load(),{},'en')
        self.assertIn('data:image/png;base64,',html);self.assertIn('&lt;script&gt;',html);self.assertNotIn('<script>alert',html)
    def test_input_formula_rejected_but_calculation_caches_ignored(self):
        self.wb['Main']['D5']=99999
        self.assertEqual(self.load()[0]['seats']['A-1']['Fare BDT'],300)
        self.wb['Main']['B5']='="DIRECT"'
        with self.assertRaisesRegex(data.InputError,'formula'):self.load()
    def test_shipped_dropdowns_and_blank_counters(self):
        from openpyxl import load_workbook
        wb=load_workbook(app.BASE/'challan_input.xlsx')
        try:
            self.assertEqual(len(wb['Main'].data_validations.dataValidation),1)
            self.assertIn('Counters',wb['Main'].data_validations.dataValidation[0].formula1)
            self.assertIn('Routes',wb['Trip Details'].data_validations.dataValidation[0].formula1)
            self.assertIsNone(wb['Counters']['A5'].value)
            self.assertEqual(wb['Routes']['D5'].value,300)
        finally:wb.close()

if __name__=='__main__':unittest.main()
